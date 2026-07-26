#!/usr/bin/env python3
from pathlib import Path
from openpyxl import load_workbook
import sys

def dump(pathstr: str):
    p = Path(pathstr)
    wb = load_workbook(p, data_only=True, read_only=True)
    ws = wb.active
    markers = []
    for row in ws.iter_rows(min_row=1, max_row=18, max_col=28, values_only=True):
        for v in row:
            if v is not None:
                markers.append(str(v).upper())
    print(' || '.join(markers))
    wb.close()

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Usage: dump_markers.py PATH')
        sys.exit(2)
    dump(sys.argv[1])
