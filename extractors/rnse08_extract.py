#!/usr/bin/env python3
"""
rnse08_extract.py — extract an RNSE-08 (ENTP rig 188, RNSE wells) Daily
Workover Report from a Word file (.doc or .docx) into the standard dict
shape.

Source layout
-------------
Microsoft Word document.  Two tables:

  Table 0 — header + mud panel:
    r0  Puits | RNSE-08 | Appareil | | TP 188 | Rapport N° | 02 | Rapport du | 21/05/2026
    r1  Nature du puits | PPH | | But de Work Over | <objective text>
    r6  Boue | DIESEL | OBM
    r8  Surf Vol (m³) | 156 | Den | FV | HPHT Filtrate | PV | AV | Yield | Gel0 | Gel10   (labels)
    r9  Vol Puits (m³) | 62 | 0.87 | 32 | 16 | 5 | 8 | 5 | 2 | 4                          (values)
    r10 Left in Hole | | Pb | Hgs | Lgs | Ca | Electrique stability | Solide              (labels)
    r11 Dumped | | | 0.73 | 0.2 | 2 | 1420 | 1                                            (values)
    r12 Tripping | | H/E | Huile | Eau | HPHT | PH | Sel | Cake                           (labels)
    r13 Tripping | | 93/07 | 92 | 07 | / | / | / | /                                      (values)
    r14 Produits à boue | | Barite | Calcium carbonate | chaux | Sodium Chloride | UD MUL | ...  (labels)
    r15 Vol Produits (T) | | 30 | 25 | 6 | 12 | 3.040 | 3.230 | 4.125 | 2.544             (values)

  Table 1 — operations:
    r0  "Déroulement des opérations"
    r1  Timing | Tarif | Opérations           (sub-headers; Timing spans 2 cols)
    r2+ HH:MM | HH:MM | T<n> | <description>
    ... a trailing block holds "NB:" notes and an after-midnight row.

Per-op tarif codes (T1/T2/...) are present directly in the table, so no
back-assignment is needed — though the universal normalize pass in
parse_source still cleans any multiplier prefixes.

Two supervisor names appear as paragraphs after "Superviseur :" →
both go to supervisor / superintendent per the two-supervisor convention.

Distinguishing markers (used by helpers.parse_source._detect_format_word):
    - "Déroulement des opérations"
    - well "RNSE-" prefix or rig "TP 188" / "TP-188"
"""
from __future__ import annotations
import re
from datetime import datetime, time, date as date_type, timedelta
from io import BytesIO
from pathlib import Path
from typing import Union, List, Optional

from docx import Document

# Re-use the LibreOffice .doc→.docx converter and a couple of helpers
# from the TP-186 extractor so we don't duplicate that logic.
from extractors.tp186_extract import _ensure_docx, _row_cells


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _clean(v) -> str:
    if v is None: return ""
    s = str(v).replace("\u202f", " ").replace("\u00a0", " ")
    return re.sub(r"\s+", " ", s).strip()


def _float(v, default=0.0) -> float:
    if v is None: return default
    if isinstance(v, (int, float)): return float(v)
    s = _clean(v).replace(",", ".").replace(" ", "")
    if s in ("", "-", "/", "None"): return default
    m = re.match(r"^([-+]?\d+(?:\.\d+)?)", s)
    if m:
        try: return float(m.group(1))
        except ValueError: pass
    return default


def _int(v, default=0) -> int:
    return int(_float(v, float(default)))


def _date_parse(v):
    if v is None: return None
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date_type): return v
    s = _clean(v)
    m = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", s)
    if m:
        try: return date_type(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError: return None
    return None


def _time_parse(v) -> Optional[time]:
    """Parse 'HH:MM' (with optional no-break spaces) → time. '24:00' → 00:00."""
    s = _clean(v)
    if not s: return None
    m = re.match(r"^(\d{1,2})\s*[:hH]\s*(\d{2})$", s)
    if not m: return None
    h, mn = int(m.group(1)), int(m.group(2))
    return time(h % 24, mn)


def _hours_between(start_t: time, end_t: time) -> float:
    sm = start_t.hour * 60 + start_t.minute
    em = end_t.hour * 60 + end_t.minute
    if em == sm:    return 24.0
    if em > sm:     return (em - sm) / 60.0
    return (em + 1440 - sm) / 60.0


def _normalize_rig(name: str) -> str:
    s = _clean(name).upper()
    m = re.search(r"TP[\s#\-]*(\d+)", s)
    if m: return f"TP-{m.group(1)}"
    return _clean(name)


def _find_value_after(cells: List[str], label: str) -> str:
    """Exact (normalized) label match → next NON-EMPTY cell; falls back to
    startswith.  Skips blank cells between a label and its value (some
    templates leave a spacer column, e.g. 'Appareil' | '' | 'TP 188')."""
    norm = lambda s: re.sub(r"\s+", " ", s.strip().lower())
    target = norm(label)
    def _next_nonempty(start_idx):
        for j in range(start_idx, len(cells)):
            if _clean(cells[j]):
                return cells[j].strip()
        return ""
    for i, c in enumerate(cells):
        if norm(c) == target:
            return _next_nonempty(i + 1)
    for i, c in enumerate(cells):
        if norm(c).startswith(target):
            return _next_nonempty(i + 1)
    return ""


# ---------------------------------------------------------------------------
# Main extractor
# ---------------------------------------------------------------------------
def parse_rnse08(source: Union[Path, str, BytesIO]) -> dict:
    if isinstance(source, BytesIO):
        try:
            doc = Document(source)
        except Exception:
            import tempfile
            with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tf:
                tf.write(source.getvalue())
                tmp_path = Path(tf.name)
            doc = Document(str(_ensure_docx(tmp_path)))
    else:
        doc = Document(str(_ensure_docx(Path(source))))

    paragraphs = [_clean(p.text) for p in doc.paragraphs]
    tables = doc.tables

    header = {}
    mud_checks = {}
    mud_volume = {}
    chemicals = []
    text_sections = {}

    # =====================================================================
    # TABLE 0 — header + mud
    # =====================================================================
    if tables:
        t0 = tables[0]
        rows = [_row_cells(r) for r in t0.rows]

        # r0 — well / rig / report no / date
        if len(rows) > 0:
            r0 = rows[0]
            header["well_name"] = _find_value_after(r0, "puits").replace(" ", "")
            app = _find_value_after(r0, "appareil")
            if app:
                header["rig_name"] = _normalize_rig(app)
            rn = _find_value_after(r0, "rapport n°") or _find_value_after(r0, "rapport n")
            if rn:
                d = _int(rn, 0)
                if d > 0: header["day_number"] = d
            dt = _find_value_after(r0, "rapport du")
            if dt:
                d = _date_parse(dt)
                if d: header["date"] = d

        # r1 — nature + objective
        if len(rows) > 1:
            r1 = rows[1]
            nat = _find_value_after(r1, "nature du puits")
            if nat: header["well_nature"] = nat
            obj = _find_value_after(r1, "but de work over")
            if obj: header["well_objective"] = obj

        # r6 — mud type (Boue | DIESEL | OBM)
        for row in rows:
            if row and _clean(row[0]).upper().startswith("BOUE"):
                # collect non-empty values after the label
                vals = [c for c in row[1:] if _clean(c)]
                if vals:
                    mud_checks["mud_type"] = " ".join(_clean(v) for v in vals)
                break

        # Mud panel — label row r8 paired with value row r9
        _pair_mud_row(rows, "surf vol", header, mud_checks, mud_volume,
                      label_map={
                          "den": ("density", mud_checks),
                          "fv": ("fun_vis", mud_checks),
                          "hpht filtrate": ("hpht_fl", mud_checks),
                          "pv": ("pv", mud_checks),
                          "av": ("av", mud_checks),
                          "yield": ("yp", mud_checks),
                          "gel0": ("gel10sec", mud_checks),
                          "gel10": ("gel10m", mud_checks),
                      })

        # Volumes from the Surf Vol / Vol Puits rows
        for row in rows:
            c0 = _clean(row[0]).lower() if row else ""
            if c0.startswith("surf vol") and len(row) > 1:
                mud_volume["pits_volume"] = _float(row[1])
            elif c0.startswith("vol puits") and len(row) > 1:
                mud_volume["string_volume"] = _float(row[1])
        if mud_volume:
            mud_volume["total_volume"] = sum(
                v for v in (mud_volume.get("pits_volume"),
                            mud_volume.get("string_volume"))
                if isinstance(v, (int, float))
            )

        # Second mud-property block (labels r10 / values r11):
        # Pb | Hgs | Lgs | Ca | Electrique stability | Solide
        _pair_mud_row(rows, "left in hole", header, mud_checks, mud_volume,
                      label_map={
                          "pb": ("pf", mud_checks),
                          "lgs": ("lgs", mud_checks),
                          "hgs": ("hgs", mud_checks),
                          "ca": ("ca", mud_checks),
                          "éléctrique stability": ("es", mud_checks),
                          "electrique stability": ("es", mud_checks),
                          "solide": ("solid", mud_checks),
                      })

        # Third block (labels r12 / values r13):
        # H/E | Huile | Eau | HPHT | PH | Sel | Cake
        _pair_mud_row(rows, "tripping", header, mud_checks, mud_volume,
                      label_map={
                          "h/e": ("oil_water_ratio", mud_checks),
                          "huile": ("oil", mud_checks),
                          "eau": ("h2o", mud_checks),
                          "ph": ("ph", mud_checks),
                      }, ratio_keys={"oil_water_ratio"})

        # Chemicals (labels r14 / values r15)
        chem_labels = None
        chem_values = None
        for ri, row in enumerate(rows):
            c0 = _clean(row[0]).lower() if row else ""
            if c0.startswith("produits à boue") or c0.startswith("produits a boue"):
                chem_labels = row[2:]                  # skip label + blank
                if ri + 1 < len(rows):
                    chem_values = rows[ri + 1][2:]
                break
        if chem_labels:
            for i, name in enumerate(chem_labels):
                name = _clean(name)
                if not name: continue
                used = ""
                if chem_values and i < len(chem_values):
                    used = _clean(chem_values[i])
                chemicals.append({
                    "item": name, "units": "T",
                    "received": "", "used": used, "on_loc": "",
                })

    # =====================================================================
    # TABLE 1 — operations
    # =====================================================================
    activities = []
    after_midnight_parts = []
    nb_notes = []
    if len(tables) > 1:
        t1 = tables[1]
        seen_header = False
        for row in t1.rows:
            cells = _row_cells(row)
            if not cells: continue
            joined = " ".join(_clean(c) for c in cells).lower()
            if "timing" in joined and "tarif" in joined:
                seen_header = True
                continue
            if not seen_header:
                continue
            # Expect [start, end, tarif, description]
            if len(cells) < 4:
                continue
            start_raw, end_raw, tarif, desc = cells[0], cells[1], cells[2], cells[3]
            start_t = _time_parse(start_raw)
            end_t   = _time_parse(end_raw)
            desc_c  = _clean(desc)

            if start_t is not None and end_t is not None:
                # A timed op.  If it has no tarif, it's the after-midnight
                # continuation (next day's first hours) — capture separately.
                if not _clean(tarif):
                    if desc_c:
                        after_midnight_parts.append(desc_c)
                    continue
                activities.append({
                    "start_time": start_t,
                    "end_time":   end_t,
                    "hours":      _hours_between(start_t, end_t),
                    "phase_name": "",
                    "code": "", "sub": "",
                    "description": desc_c,
                    "start_md": 0, "end_md": 0,
                    "npt": 0, "npt_detail": "",
                    "npt_company": "", "op_company": "",
                    "bill": _clean(tarif).upper(),
                })
            elif desc_c:
                # No timing — a note row ("NB: ...") or continuation
                if desc_c.upper().startswith("NB"):
                    nb_notes.append(desc_c)
                elif activities:
                    activities[-1]["description"] = (
                        activities[-1]["description"] + "\n" + desc_c
                    ).strip()

    # =====================================================================
    # TARIF TOTALS — derived by summing per-op hours per code
    # =====================================================================
    tarif_totals = {}
    for a in activities:
        code = _clean(a.get("bill")).lower()
        if re.match(r"^t\d+$", code):
            tarif_totals[code] = tarif_totals.get(code, 0.0) + a.get("hours", 0.0)

    # =====================================================================
    # SUPERVISORS — paragraphs after "Superviseur :"
    # Two names → supervisor + superintendent (two-supervisor convention).
    # =====================================================================
    sup_names = []
    grab = False
    for p in paragraphs:
        if not p: continue
        if "SUPERVISEUR" in p.upper():
            grab = True
            # Name might be on the same line after the colon
            after = re.sub(r".*superviseur\s*:\s*", "", p, flags=re.IGNORECASE).strip()
            if after:
                sup_names.append(after)
            continue
        if grab and len(sup_names) < 2 and len(p) < 50:
            # Looks like a name line (short, has a dot-initial or caps)
            if re.match(r"^[A-Z]\.?\s*[A-ZÉÈ]", p) or "." in p[:4]:
                sup_names.append(p)
    if len(sup_names) >= 1:
        header["supervisor"] = sup_names[0]
    if len(sup_names) >= 2:
        header["superintendent"] = sup_names[1]

    # =====================================================================
    # TEXT SECTIONS
    # =====================================================================
    if after_midnight_parts:
        text_sections["after_midnight"] = " ".join(after_midnight_parts)
    # Situation = last op description (capped at 300 chars)
    if activities:
        sit = activities[-1]["description"]
        if len(sit) > 300:
            sit = sit[:300].rsplit(None, 1)[0] + "…"
        text_sections["current_operation"] = sit
        text_sections["day_summary"] = sit
    if nb_notes:
        text_sections["remarks"] = " ".join(nb_notes)

    return {
        "header": header,
        "activities": activities,
        "text_sections": text_sections,
        "mud_checks": mud_checks,
        "mud_volume": mud_volume,
        "mud_chemical_usage": chemicals,
        "personnel_data": [],
        "pumps": [],
        "well_location": {},
        "survey_data": [],
        "safety": {},
        "tarif_totals": tarif_totals,
    }


def _pair_mud_row(rows, label_prefix, header, mud_checks, mud_volume,
                  label_map, ratio_keys=None):
    """Find a label row whose first cell starts with label_prefix, pair it
    with the row directly below (the value row), and map labels→values
    using label_map { normalized_label: (dest_key, dest_dict) }."""
    ratio_keys = ratio_keys or set()
    norm = lambda s: re.sub(r"\s+", " ", _clean(s).lower())
    for ri, row in enumerate(rows):
        if not row: continue
        if norm(row[0]).startswith(norm(label_prefix)):
            if ri + 1 >= len(rows):
                return
            label_row = row
            value_row = rows[ri + 1]
            for ci, lab in enumerate(label_row):
                key = norm(lab)
                if key in label_map and ci < len(value_row):
                    dest_key, dest_dict = label_map[key]
                    raw = value_row[ci]
                    if dest_key in ratio_keys:
                        v = _clean(raw)
                        if v and v != "/":
                            dest_dict[dest_key] = v
                    else:
                        fv = _float(raw)
                        if fv or fv == 0.0:
                            # Only store if the cell actually had a number
                            if _clean(raw) not in ("", "/"):
                                dest_dict[dest_key] = fv
            return


# Drop-in compat
parse_daily_excel_report = parse_rnse08


if __name__ == "__main__":
    import sys, json
    if len(sys.argv) < 2:
        sys.exit("Usage: rnse08_extract.py SOURCE.doc|SOURCE.docx")
    data = parse_rnse08(Path(sys.argv[1]))

    def default(o):
        if isinstance(o, (date_type, datetime)): return o.isoformat()
        if isinstance(o, time): return o.strftime("%H:%M:%S")
        if isinstance(o, timedelta): return o.total_seconds()
        return str(o)
    print(json.dumps(data, indent=2, default=default, ensure_ascii=False))
