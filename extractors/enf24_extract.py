"""
enf24_extract.py — Extract structured data from the newer ENAFOR Daily Drilling
Report template (observed on rig ENF#24, sheet "DDR 06-07-2026").

This is a sibling of enf17_extract.py / parse_ddr(). Most of the sheet layout
is identical to the ENF#17 template (personnel grid, safety counters, mud
checks, mud volume, BOP test date, activities table) but a handful of header
cells moved by one row/column, and the "Last Csg" block was restructured from
two columns (size, depth) into a single value column keyed by row label.
See the "DIFFERENCES FROM enf17_extract.py" comment block below for the full
list of what changed and why.

Output shape is unchanged from parse_ddr() so this remains a drop-in for
insert_parsed_report() / excel_import_service.parse_daily_excel_report().

CLI:
    python enf24_extract.py SOURCE.xlsx [-o OUTPUT.json] [--pretty]
"""

from __future__ import annotations
import argparse
import json
import re
import sys
from datetime import datetime, time, date as date_type, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl


# ===========================================================================
# Generic helpers (identical to enf17_extract.py — unchanged parsing rules)
# ===========================================================================

def _clean(text) -> str:
    if text is None:
        return ""
    return re.sub(r"\s+", " ", str(text)).strip()


def _float(text, default: float = 0.0) -> float:
    if text is None:
        return default
    if isinstance(text, (int, float)):
        return float(text)
    s = _clean(str(text)).replace(",", ".").replace(" ", "")
    if s in ("", "-", "/", "()", "None"):
        return default
    s = re.sub(r"[a-zA-Z%]+$", "", s).strip()
    if not s:
        return default
    try:
        return float(s)
    except (ValueError, TypeError):
        return default


def _int(text, default: int = 0) -> int:
    return int(_float(text, float(default)))


def _date_parse(val) -> Optional[date_type]:
    """Parse a date from many representations: datetime, date, 'dd/mm/yyyy',
    'dd-mm-yyyy', and the ENAFOR-specific 'dd,mm,yy' (e.g. '24,04,26')."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date_type):
        return val
    text = _clean(str(val))
    m = re.match(r"^\s*(\d{1,2})[,\.\-/](\d{1,2})[,\.\-/](\d{2,4})\s*$", text)
    if m:
        d, mo, y = (int(x) for x in m.groups())
        if y < 100:
            y += 2000
        try:
            return date_type(y, mo, d)
        except ValueError:
            pass
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d %m %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _time_from_cell(val) -> Optional[time]:
    if val is None:
        return None
    if isinstance(val, time):
        return val
    if isinstance(val, datetime):
        return val.time()
    if isinstance(val, timedelta):
        secs = int(val.total_seconds())
        if secs < 0:
            return None
        return time((secs // 3600) % 24, (secs % 3600) // 60)
    text = _clean(str(val))
    if text in ("", "-", "None"):
        return None
    if text == "24:00":
        return time(0, 0)
    for fmt in ("%H:%M:%S", "%H:%M", "%Hh%M", "%H.%M"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    return None


def _hours_from_cell(val) -> float:
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, timedelta):
        return val.total_seconds() / 3600
    if isinstance(val, time):
        return val.hour + val.minute / 60.0
    if isinstance(val, datetime):
        return val.hour + val.minute / 60.0
    return _float(val)


def _build_merged_lookup(ws) -> Dict[Tuple[int, int], Any]:
    lookup: Dict[Tuple[int, int], Any] = {}
    for mr in ws.merged_cells.ranges:
        val = ws.cell(row=mr.min_row, column=mr.min_col).value
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                if r != mr.min_row or c != mr.min_col:
                    lookup[(r, c)] = val
    return lookup


def _cell(ws, row: int, col: int, lookup: Dict) -> Any:
    val = ws.cell(row=row, column=col).value
    if val is not None:
        return val
    return lookup.get((row, col))


def _scan_for(ws, lookup, keyword, row_range, col_range):
    kw = keyword.upper()
    for r in range(row_range[0], row_range[1] + 1):
        for c in range(col_range[0], col_range[1] + 1):
            v = _cell(ws, r, c, lookup)
            if v and kw in str(v).upper():
                return (r, c)
    return None


BBL_TO_M3 = 0.158987294928


# ===========================================================================
# DIFFERENCES FROM enf17_extract.py (found by diffing this file's cells
# against the enf17_extract.py coordinate assumptions):
#
# 1. Present depth: the ENF17 template had the value directly in row 5
#    (J5). In this template row 5 is a pure label row ("Present depth (m)"
#    spans J5:L5) and the VALUE lives one row down, in J6 (e.g. '3437m').
#
# 2. Drill-time / Overall-days swapped a column: ENF17 read drill-time
#    hours from S6. In this template S5/S6 is actually labelled
#    "Overall-days" (col 19) and "Drill-time (H)" is one column to the
#    left, labelled at R5 with its value (when filled in) at R6 (col 18).
#    We now capture BOTH: drill_time_hours (R6) and the new overall_days
#    (S6) field.
#
# 3. "Last Csg" block was reshuffled from a size+depth pair per row
#    (D11/E11, D12/E12, D13) into a single value column (col D) keyed by
#    the row label in col C: row11 label "Size" -> D11 (casing size),
#    row12 label "Top-shoe" -> D12 (shoe depth), row13 label "Top-Liner"
#    -> D13 (liner depth). There is no separate "size" for top-shoe/liner
#    in this template.
#
# 4. No "Report N°" label exists anywhere on this sheet (searched
#    exhaustively). The closest analogue is "Overall-days" (S6), a running
#    day count, which we use as day_number when present. The old N°-label
#    scan is kept as a fallback in case a future report reintroduces it.
#
# 5. Narrative text after the ops table ("AFTER MIDNIGHT" section onward)
#    is looser free text spanning several paragraph blocks rather than a
#    single labelled value next to the marker. We now walk forward from
#    the marker, grouping consecutive non-blank rows into paragraphs: the
#    first paragraph is stored as after_midnight, later paragraphs are
#    joined into additional_remarks. The "Plan Opetations :" label at the
#    end of that block had no adjacent value in this sample report — if a
#    future report fills it in (same row, cols to the right), it will
#    still be captured via the existing PLAN scan.
#
# Personnel grid (rows 14-17), safety counters (rows 7-8), mud checks
# (rows 26-30, cols R-U), mud volume (rows 24 & 35), BOP test date (U9),
# supervisor/superintendent (G3/G4), field/well (G5/G6), rig name (A3),
# report date (W4) and the activities table (rows 20+, cols A-E) are all
# in the SAME cells as enf17_extract.py — verified against this file.
# ===========================================================================


def _extract_header(ws, lookup) -> Dict[str, Any]:
    h: Dict[str, Any] = {}

    rig = _cell(ws, 3, 1, lookup)
    if rig:
        h["rig_name"] = re.sub(r"\s*#\s*", "#", _clean(str(rig)))

    d = _date_parse(_cell(ws, 4, 23, lookup))          # W4
    if d:
        h["date"] = d

    field = _cell(ws, 5, 7, lookup)                    # G5
    if field:
        h["field_name"] = _clean(str(field))

    well = _cell(ws, 6, 7, lookup)                      # G6
    if well:
        h["well_name"] = _clean(str(well)).replace("# ", "#").replace(" #", "#")

    # Previous depth: M6 (unchanged position vs enf17_extract.py)
    pdepth = _cell(ws, 6, 13, lookup)
    if pdepth is not None:
        h["previous_depth_text"] = _clean(str(pdepth))
        h["tmd"] = _float(pdepth)

    # Present depth: MOVED from J5 to J6 in this template (see note #1)
    present = _cell(ws, 6, 10, lookup)                  # J6
    if present is not None and _float(present) > 0:
        h["present_depth"] = _float(present)

    # Drill time (h): MOVED from S6 to R6 in this template (see note #2)
    dt = _cell(ws, 6, 18, lookup)                       # R6
    if dt is not None:
        h["drill_time_hours"] = _float(dt)

    # New: overall running day count, col S6 (see note #2 / #4)
    od = _cell(ws, 6, 19, lookup)                       # S6
    if od is not None and _clean(str(od)) != "":
        h["overall_days"] = _int(od)

    # Report N°: no dedicated label found on this template (see note #4).
    # Try the legacy scan first (harmless no-op if absent), else fall back
    # to overall_days as the closest available "day count" analogue.
    pos = _scan_for(ws, lookup, "N°", (1, 8), (20, 28))
    found_day_number = False
    if pos:
        for dc in (1, 2, 3):
            n = _cell(ws, pos[0], pos[1] + dc, lookup)
            if n is None or isinstance(n, (datetime, date_type, time)):
                continue
            txt = _clean(str(n))
            if txt in ("", "N°"):
                continue
            digits = re.sub(r"\D", "", txt)
            if digits and len(digits) < 6:
                h["day_number"] = int(digits)
                found_day_number = True
                break
    if not found_day_number and "overall_days" in h:
        h["day_number"] = h["overall_days"]

    # BOP test date: U9 (unchanged)
    bop = _date_parse(_cell(ws, 9, 21, lookup))
    if bop:
        h["bop_test"] = bop

    # TP Sénior (G3) -> supervisor; TP Junior (G4) -> superintendent
    # (unchanged mapping from enf17_extract.py)
    ts = _cell(ws, 3, 7, lookup)
    if ts:
        h["supervisor"] = _clean(str(ts))
    tj = _cell(ws, 4, 7, lookup)
    if tj:
        h["superintendent"] = _clean(str(tj))

    # Last Csg block: RESTRUCTURED in this template (see note #3).
    # Row 11 label "Size" -> D11 = casing size string (e.g. "7''")
    # Row 12 label "Top-shoe" -> D12 = shoe depth (m)
    # Row 13 label "Top-Liner" -> D13 = liner depth (m)
    csg_size = _cell(ws, 11, 4, lookup)      # D11
    shoe_depth = _cell(ws, 12, 4, lookup)    # D12
    liner_depth = _cell(ws, 13, 4, lookup)   # D13

    if csg_size:
        h["last_csg_size"] = _clean(str(csg_size))
    if shoe_depth is not None and _clean(str(shoe_depth)) != "":
        h["top_shoe_depth"] = _float(shoe_depth)
        if csg_size:
            h["last_csg_shoe"] = f'{_clean(str(csg_size))} @ {shoe_depth}m'
    if liner_depth is not None and _clean(str(liner_depth)) != "":
        h["top_liner_depth"] = _float(liner_depth)

    # BHA: G11 (length) + J11 (details) — unchanged positions, both blank
    # in the sample report but kept for when they're filled in.
    bha_len = _cell(ws, 11, 7, lookup)
    if bha_len:
        h["bha_length"] = _clean(str(bha_len))
    bha_det = _cell(ws, 11, 10, lookup)
    if bha_det:
        h["bha_details"] = _clean(str(bha_det))

    return h


def _extract_activities(ws, lookup) -> List[Dict[str, Any]]:
    """Operations / timing block — same cells/logic as enf17_extract.py."""
    ops: List[Dict[str, Any]] = []
    current: Optional[Dict[str, Any]] = None

    STOP_MARKERS = ("AFTER MIDNIGHT", "APRÈS MINUIT", "APRES MINUIT",
                    "PLAN OPER", "PROGRAMME", "SITUATION", "REMARKS",
                    "REMARQUES", "PERSONNEL", "MUD CHECK")

    OPS_START = 20
    OPS_MAX_END = 50

    for r in range(OPS_START, OPS_MAX_END):
        start_val = _cell(ws, r, 1, lookup)
        end_val   = _cell(ws, r, 2, lookup)
        hours_val = _cell(ws, r, 3, lookup)
        code_val  = _cell(ws, r, 4, lookup)
        desc_val  = _cell(ws, r, 5, lookup)

        row_marker_text = " ".join(
            str(_cell(ws, r, c, lookup) or "") for c in range(1, 6)
        ).upper()
        if any(m in row_marker_text for m in STOP_MARKERS):
            break

        start_t = _time_from_cell(start_val)
        end_t   = _time_from_cell(end_val)

        if start_t is not None and end_t is not None:
            if current:
                ops.append(current)
            current = {
                "start_time":   start_t,
                "end_time":     end_t,
                "hours":        _hours_from_cell(hours_val),
                "phase_name":   "",
                "code":         "",
                "sub":          "",
                "description":  _clean(str(desc_val or "")),
                "start_md":     0,
                "end_md":       0,
                "npt":          0,
                "npt_detail":   "",
                "npt_company":  "",
                "op_company":   "",
                "bill":         _clean(str(code_val or "")),
            }
        elif desc_val and current is not None:
            extra = _clean(str(desc_val))
            current["description"] = (current["description"] + "\n" + extra).strip()
    if current:
        ops.append(current)
    return ops


def _extract_text_sections(ws, lookup, activities) -> Dict[str, str]:
    """Narrative text blocks. Adapted for this template's looser paragraph
    layout after the ops table (see note #5)."""
    sec: Dict[str, str] = {}

    pos = _scan_for(ws, lookup, "AFTER MIDNIGHT", (25, 60), (1, 8))
    if not pos:
        pos = _scan_for(ws, lookup, "APRÈS MINUIT", (25, 60), (1, 8))
    if not pos:
        pos = _scan_for(ws, lookup, "APRES MINUIT", (25, 60), (1, 8))

    if pos:
        marker_row = pos[0]
        paragraphs: List[List[str]] = []
        current_para: List[str] = []
        r = marker_row + 1
        max_r = marker_row + 40
        while r <= max_r:
            row_text_all = " ".join(
                str(_cell(ws, r, c, lookup) or "") for c in range(1, 6)
            ).upper()
            # Stop once we reach the "Daily cost" footer row that follows
            # this whole narrative block.
            if "DAILY COST" in row_text_all:
                break
            v = _cell(ws, r, 5, lookup)  # col E
            text = _clean(str(v)) if v is not None else ""
            if text:
                current_para.append(text)
            else:
                if current_para:
                    paragraphs.append(current_para)
                    current_para = []
            r += 1
        if current_para:
            paragraphs.append(current_para)

        if paragraphs:
            first = "\n".join(paragraphs[0])
            first = re.sub(r"^\s*(?:AFTER\s*MIDNIGHT|APR[ÈE]S\s*MINUIT)\s*:?\s*",
                            "", first, flags=re.IGNORECASE)
            if first:
                sec["after_midnight"] = first
            if len(paragraphs) > 1:
                rest = "\n\n".join("\n".join(p) for p in paragraphs[1:])
                if rest:
                    sec["additional_remarks"] = rest

    # Plan operations: same-row scan to the right of the label, as before.
    # In the sample report this label had no adjacent value (left blank).
    pos = _scan_for(ws, lookup, "PLAN", (40, 70), (1, 12))
    if pos:
        for dc in (1, 2, 3):
            v = _cell(ws, pos[0], pos[1] + dc, lookup)
            if v and "PLAN" not in str(v).upper():
                sec["plan_operations"] = _clean(str(v))
                break

    SITUATION_MAX = 300
    if activities:
        in_prog = [o for o in activities if "IN PROGRESS" in (o.get("description") or "").upper()]
        chosen = in_prog[-1] if in_prog else activities[-1]
        text = chosen["description"] or ""
        if len(text) > SITUATION_MAX:
            cut = text[:SITUATION_MAX].rsplit(None, 1)[0]
            text = cut + "…"
        sec["current_operation"] = text
        sec["day_summary"] = text

    return sec


def _extract_mud_checks(ws, lookup) -> Dict[str, Any]:
    """Mud properties — same cells as enf17_extract.py (verified)."""
    mud: Dict[str, Any] = {}

    def _put(key, row, col):
        v = _cell(ws, row, col, lookup)
        if v is not None and (isinstance(v, (int, float)) or _clean(str(v)) not in ("", "-")):
            mud[key] = _float(v) if not isinstance(v, str) or re.search(r"\d", v) else _clean(str(v))

    _put("density",   26, 19)   # S26
    _put("fun_vis",   26, 21)   # U26
    v = _cell(ws, 27, 19, lookup)
    if v is not None:
        mud["solid"] = _float(v)
    _put("yp",        27, 21)   # U27
    _put("gel10sec",  28, 19)   # S28
    _put("h2o",       29, 21)   # U29
    _put("lgs",       30, 21)   # U30

    ow = _cell(ws, 30, 19, lookup)
    if ow:
        mud["oil_water_ratio"] = _clean(str(ow))

    pos = _scan_for(ws, lookup, "MUD", (5, 12), (15, 25))
    if pos:
        v = _cell(ws, pos[0], pos[1] + 1, lookup)
        if v:
            mud["mud_type"] = _clean(str(v))

    return mud


def _extract_mud_volume(ws, lookup) -> Dict[str, Any]:
    """Hole volume / circulation volume — same cells as enf17_extract.py."""
    vol: Dict[str, Any] = {}

    hole_bbl = _cell(ws, 24, 19, lookup)        # S24
    if hole_bbl is not None:
        vol["string_volume_bbl"] = _float(hole_bbl)
        vol["string_volume"]     = round(_float(hole_bbl) * BBL_TO_M3, 2)

    circ_bbl = _cell(ws, 24, 21, lookup)        # U24
    if circ_bbl is not None:
        vol["total_volume_bbl"]  = _float(circ_bbl)
        vol["total_volume"]      = round(_float(circ_bbl) * BBL_TO_M3, 2)

    rw = _cell(ws, 35, 20, lookup)              # T35
    if rw is not None:
        vol["rig_water_m3"] = _float(rw)

    # New in this template: loss volumes appear next to the hole/circ vol
    # labels (row 25) — not present in enf17_extract.py, added here as they
    # were clearly visible in the sample sheet.
    loss_trip = _cell(ws, 25, 19, lookup)       # S25
    if loss_trip is not None and _clean(str(loss_trip)) != "":
        vol["loss_tripping_vol_bbl"] = _float(loss_trip)
    loss_hole = _cell(ws, 25, 21, lookup)       # U25
    if loss_hole is not None and _clean(str(loss_hole)) != "":
        vol["loss_hole_vol_bbl"] = _float(loss_hole)

    return vol


def _extract_personnel(ws, lookup) -> List[Dict[str, Any]]:
    """Crew counts — same rows/columns as enf17_extract.py (verified)."""
    rows = []
    mapping = [
        ("S/Tool Pusher",            15, 3),
        ("J/Tool Pusher",            15, 5),
        ("Driller",                  15, 7),
        ("A/Driller",                15, 10),
        ("Cont. Lab",                15, 12),
        ("Cas. Lab",                 15, 15),
        ("Militaires/Vigiles",       15, 17),
        ("SPV (Elec/Mec/TDS)",       15, 18),
        ("DTR/Personnels",           15, 19),
        ("Front loader/Cariste",     15, 21),
        ("Intendant",                15, 25),
        ("Elect",                    17, 1),
        ("Mechanic",                 17, 3),
        ("Driver",                   17, 5),
        ("Catering",                 17, 7),
        ("Company man",              17, 9),
        ("Trainee",                  17, 12),
        ("HSE",                      17, 15),
        ("ADG",                      17, 17),
        ("Medec/Infirmier",          17, 18),
        ("Welder",                   17, 19),
        ("Other",                    17, 21),
    ]
    for label, r, c in mapping:
        v = _cell(ws, r, c, lookup)
        if v is None or _clean(str(v)) == "":
            continue
        text = _clean(str(v))
        if re.match(r"^\d+(\.\d+)?$", text):
            n = _int(text)
        elif re.match(r"^[\d.]+(\s*\+\s*[\d.]+)+$", text):
            n = sum(_int(p.strip()) for p in text.split("+"))
        elif re.match(r"^[\d.]+(/[\d.]+)+$", text):
            n = sum(_int(p.strip()) for p in text.split("/"))
        else:
            n = 0
        names_field = "" if re.match(r"^\d+(\.\d+)?$", text) else text
        rows.append({"company": label, "number": n, "hours": "", "names": names_field})
    return rows


def _extract_safety(ws, lookup) -> Dict[str, Any]:
    """HSE counters — same cells as enf17_extract.py (verified)."""
    safety: Dict[str, Any] = {}
    afd = _cell(ws, 7, 6, lookup)               # F7
    if afd is not None:
        safety["accident_free_days"] = _int(afd)
    m = _cell(ws, 7, 21, lookup)                # U7
    if m is not None:
        safety["hse_meetings"] = _clean(str(m))
    p = _cell(ws, 8, 21, lookup)                # U8
    if p is not None:
        safety["permits_to_work"] = _clean(str(p))
    s = _cell(ws, 7, 25, lookup)                # Y7
    if s is not None:
        safety["stop_cards"] = _clean(str(s))
    e = _cell(ws, 8, 25, lookup)                # Y8
    if e is not None:
        safety["exercises"] = _clean(str(e))
    return safety


def _extract_npt_summary(ws, lookup) -> Dict[str, Any]:
    """NEW: NPT/T0 KPI block (rows 57-62, label cols A-D, value col E).
    Not present in enf17_extract.py — this template exposes running NPT
    totals/percentages that weren't part of the older layout."""
    npt: Dict[str, Any] = {}

    def _put(key, row):
        v = _cell(ws, row, 5, lookup)  # col E
        if v is not None and _clean(str(v)) != "":
            npt[key] = _clean(str(v))

    _put("npt_last_24h",        57)
    _put("npt_current_month",   58)
    _put("npt_current_well",    59)
    _put("t0_current_month",    60)
    _put("npt_month_pct",       61)
    _put("npt_well_pct",        62)
    return npt


def _extract_tarif_totals(activities) -> Dict[str, float]:
    tally: Dict[str, float] = {}
    for op in activities:
        code = (op.get("bill") or "").strip().upper()
        if code:
            tally[code] = tally.get(code, 0.0) + _hours_from_cell(op.get("hours"))
    return tally


# ===========================================================================
# Main entry point
# ===========================================================================

def parse_ddr(source) -> Dict[str, Any]:
    """Parse this template's Daily Drilling Report and return a structured
    dict with the same shape as enf17_extract.parse_ddr()."""
    wb = openpyxl.load_workbook(source, data_only=True)
    ws = wb.active
    lookup = _build_merged_lookup(ws)

    header     = _extract_header(ws, lookup)
    activities = _extract_activities(ws, lookup)
    text_sec   = _extract_text_sections(ws, lookup, activities)

    result: Dict[str, Any] = {
        "header":             header,
        "activities":         activities,
        "text_sections":      text_sec,
        "mud_checks":         _extract_mud_checks(ws, lookup),
        "mud_volume":         _extract_mud_volume(ws, lookup),
        "mud_chemical_usage": [],
        "personnel_data":     _extract_personnel(ws, lookup),
        "pumps":              [],
        "well_location":      {},
        "survey_data":        [],
        "safety":             _extract_safety(ws, lookup),
        "tarif_totals":       _extract_tarif_totals(activities),
        # New extra vs enf17_extract.py — see _extract_npt_summary docstring.
        "npt_summary":        _extract_npt_summary(ws, lookup),
    }

    wb.close()
    return result


def parse_daily_excel_report(file_bytes: bytes) -> Dict[str, Any]:
    """Drop-in compatibility wrapper matching excel_import_service's
    parse_daily_excel_report() signature."""
    data = parse_ddr(BytesIO(file_bytes))
    if not data["header"].get("well_name"):
        raise ValueError("Could not parse well name from Excel file")
    if not data["header"].get("date"):
        raise ValueError("Could not parse date from Excel file")
    return data


# ---------------------------------------------------------------------------
# JSON CLI
# ---------------------------------------------------------------------------

def _json_safe(obj):
    if isinstance(obj, (datetime, date_type)):
        return obj.isoformat()
    if isinstance(obj, time):
        return obj.strftime("%H:%M:%S")
    if isinstance(obj, dict):
        return {k: _json_safe(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [_json_safe(x) for x in obj]
    return obj


def main(argv=None) -> int:
    p = argparse.ArgumentParser(
        description="Extract structured data from this ENAFOR DDR template."
    )
    p.add_argument("source", type=Path, help="Path to the source .xlsx file")
    p.add_argument("-o", "--output", type=Path, default=None,
                   help="Output JSON path (default: <source>.json next to the source)")
    p.add_argument("--pretty", action="store_true",
                   help="Pretty-print the JSON output (indent=2)")
    args = p.parse_args(argv)

    if not args.source.exists():
        sys.exit(f"ERROR: source file not found: {args.source}")

    data = parse_ddr(args.source)

    if args.output is None:
        args.output = args.source.with_suffix(".json")

    with open(args.output, "w", encoding="utf-8") as fh:
        json.dump(_json_safe(data), fh,
                  indent=2 if args.pretty else None,
                  ensure_ascii=False)

    h = data["header"]
    print(f"Wrote {args.output}")
    print(f"  well={h.get('well_name')!r}  rig={h.get('rig_name')!r}  "
          f"field={h.get('field_name')!r}  date={h.get('date')}")
    print(f"  {len(data['activities'])} activities | "
          f"tarif totals: {data['tarif_totals']}")
    print(f"  {len(data['personnel_data'])} personnel rows | "
          f"{len(data['mud_checks'])} mud props | "
          f"safety: {data['safety']} | npt_summary: {data['npt_summary']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
