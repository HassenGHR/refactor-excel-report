#!/usr/bin/env python3
"""
enf08_extract.py — extract an ENF-08 (ENAFOR rig 08, ISNO wells, TINRHERT
field) Daily Workover Report from a Word file (.doc or .docx) into the
standard dict shape.

Source layout
-------------
Microsoft Word document (legacy .doc or modern .docx).  Same overall
title as the TP-186 report family — "RAPPORT JOURNALIER WORK-OVER" —
but for rig ENF#08 / well ISNO#xx / field TINRHERT, with a few real
differences from the TP-186 layout (see notes below).

The document uses 4 tables:
    Table 0 — header (well, field, rig, tubing/tool designation, mud
              params, composition, technical params)
    Table 1 — operations + chemicals + daily cost + situation + plan
              Operations live in a SINGLE table row, split across two
              PARALLEL cells: cell 0 holds newline-separated time bands
              ("HHhMM-HHhMM"), cell 1 holds the newline-separated
              descriptions aligned by line index (same convention as
              TP-186's parallel-cell layout, just one row down instead
              of two).
    Table 2 — tarif breakdown: T.1/T.2/T.3/T.4/N-R hours and amounts (DA)
    Table 3 — service companies and their costs

Differences from the TP-186 extractor that this file specifically
accounts for:
    - The "HORAIRE" header cell is misspelled "HORRAIRE" in this report
      family, so both spellings are matched.
    - Table 1 carries an explicit "Coût journalier" (daily cost) cell —
      used in preference to hunting for the table 2 TOTAL row, though
      the TOTAL row is kept as a fallback.
    - The mud "Filtrat" line in this report is a Huile/Eau (oil/water)
      split ratio (e.g. "90/10"), not a single numeric filtrate value,
      so it is captured separately rather than forced into a float.
    - Table 0 carries a tubing/completion "OUTIL" designation
      (Désignation / Diamètre / Marque-type) and a "COMPOSITION
      GARNITURE" (DP/DC) block instead of TP-186's BHA/bit fields;
      these are captured as their own header fields.
    - Rig names use the "ENF" prefix ("ENF # 08" / "ENF#08") rather
      than "TP".

.doc files are auto-converted to .docx using LibreOffice headless (must
be installed on the deployment host).  .docx files are read directly.

Distinguishing markers (used by helpers.parse_source._detect_format):
    - file extension .doc or .docx
    - contains "RAPPORT JOURNALIER WORK-OVER" together with
      "ENF # 08" / "ENF#08" / "ENF-08" (or field "TINRHERT" / well
      "ISNO")
"""
from __future__ import annotations
import re
import subprocess
import tempfile
from datetime import datetime, time, date as date_type, timedelta
from io import BytesIO
from pathlib import Path
from typing import Union, List, Tuple, Optional

from docx import Document


# ---------------------------------------------------------------------------
# Helpers (shared conventions with tp186_extract.py)
# ---------------------------------------------------------------------------
def _clean(v) -> str:
    if v is None: return ""
    s = str(v)
    # Strip Word's narrow no-break space (\u202f) and regular no-break (\u00a0)
    s = s.replace("\u202f", " ").replace("\u00a0", " ")
    return re.sub(r"\s+", " ", s).strip()


def _float(v, default=0.0) -> float:
    if v is None: return default
    if isinstance(v, (int, float)): return float(v)
    s = _clean(v).replace(",", ".").replace(" ", "")
    if s in ("", "-", "/", "None"): return default
    # Strip trailing unit (handles 'm3', 'm³', '%', 'kg', 'psi', etc.):
    # take everything up to the last digit + optional decimal portion.
    m = re.match(r"^([-+]?\d+(?:\.\d+)?)", s)
    if m:
        try: return float(m.group(1))
        except ValueError: pass
    return default


def _int(v, default=0) -> int:
    return int(_float(v, float(default)))


def _money_to_int(s) -> int:
    """'2 350 649,53 DA' or '2\\u202f350\\u202f649' → 2350649."""
    if s is None: return 0
    if isinstance(s, (int, float)): return int(s)
    cleaned = _clean(s)
    # Drop currency suffix
    cleaned = re.sub(r"\s*DA\s*$", "", cleaned, flags=re.IGNORECASE)
    # Drop decimals (we want integer DA)
    cleaned = re.split(r"[.,]\d{2,3}\s*$", cleaned)[0]
    # Strip non-digits
    digits = re.sub(r"\D", "", cleaned)
    return int(digits) if digits else 0


def _parse_hhmm_duration(s: str) -> float:
    """'22H30' / '22h30' / '00H00' → decimal hours."""
    s = _clean(s)
    m = re.match(r"^(\d{1,2})[hH:](\d{2})$", s)
    if not m: return 0.0
    return int(m.group(1)) + int(m.group(2)) / 60.0


def _hours_between(start_t: time, end_t: time) -> float:
    sm = start_t.hour * 60 + start_t.minute
    em = end_t.hour * 60 + end_t.minute
    if em == sm:    return 24.0
    if em > sm:     return (em - sm) / 60.0
    return (em + 1440 - sm) / 60.0


def _find_libreoffice() -> str:
    """Locate the LibreOffice executable across platforms.

    Search order:
      1. LIBREOFFICE_PATH environment variable (lets the user override)
      2. `soffice` / `libreoffice` on PATH (Linux / macOS / Windows-with-PATH)
      3. Common Windows install locations (Program Files / Program Files (x86))
      4. macOS .app bundle
    Returns the executable path as a string, or raises RuntimeError with
    a platform-specific install hint.
    """
    import os, shutil, sys

    # 1. Explicit override
    env = os.environ.get("LIBREOFFICE_PATH")
    if env and Path(env).exists():
        return env

    # 2. On PATH — both common command names
    for cmd in ("soffice", "libreoffice"):
        found = shutil.which(cmd)
        if found:
            return found

    # 3. Windows default install locations
    if sys.platform.startswith("win"):
        candidates = [
            Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
            Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
            Path(r"C:\Program Files\LibreOffice 7\program\soffice.exe"),
            Path(r"C:\Program Files\LibreOffice 25\program\soffice.exe"),
        ]
        for p in candidates:
            if p.exists():
                return str(p)

    # 4. macOS bundle
    if sys.platform == "darwin":
        for p in (
            Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
            Path.home() / "Applications/LibreOffice.app/Contents/MacOS/soffice",
        ):
            if p.exists():
                return str(p)

    # Not found — raise with platform-appropriate hint
    if sys.platform.startswith("win"):
        hint = (
            "LibreOffice not found. Install it from https://www.libreoffice.org/\n"
            "and either:\n"
            "  (a) add C:\\Program Files\\LibreOffice\\program to your PATH, or\n"
            "  (b) set the LIBREOFFICE_PATH environment variable to the full\n"
            "      path of soffice.exe, e.g.\n"
            "      setx LIBREOFFICE_PATH \"C:\\Program Files\\LibreOffice\\program\\soffice.exe\""
        )
    elif sys.platform == "darwin":
        hint = ("LibreOffice not found. Install via "
                "`brew install --cask libreoffice` or download from "
                "https://www.libreoffice.org/")
    else:
        hint = ("LibreOffice not found on PATH. Install with "
                "`apt-get install libreoffice` (Debian/Ubuntu) or set "
                "LIBREOFFICE_PATH to the soffice executable.")
    raise RuntimeError(hint)


def _ensure_docx(source_path: Path) -> Path:
    """If source is .doc (legacy binary), convert to .docx using LibreOffice
    headless.  Returns the path to the .docx file (may be the same path if
    the source was already .docx)."""
    suffix = source_path.suffix.lower()
    if suffix == ".docx":
        return source_path
    if suffix != ".doc":
        raise ValueError(f"Unsupported Word file extension: {suffix}")

    soffice = _find_libreoffice()

    tmp_dir = Path(tempfile.mkdtemp(prefix="docconv_"))
    try:
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "docx",
             "--outdir", str(tmp_dir), str(source_path)],
            capture_output=True, text=True, timeout=120,
        )
    except FileNotFoundError:
        raise RuntimeError(
            f"LibreOffice was located at {soffice!r} but couldn't be executed. "
            "Check that the file exists and is runnable, or set "
            "LIBREOFFICE_PATH to a working soffice executable."
        )
    except subprocess.TimeoutExpired:
        raise RuntimeError("LibreOffice conversion timed out (>120s)")
    if result.returncode != 0:
        raise RuntimeError(
            f"LibreOffice conversion failed: {result.stderr[:500]}"
        )
    docx_path = tmp_dir / (source_path.stem + ".docx")
    if not docx_path.exists():
        raise RuntimeError(
            f"Expected output not found after conversion: {docx_path}"
        )
    return docx_path


def _date_from_text(s: str) -> Optional[date_type]:
    """Extract first DD/MM/YYYY or DD-MM-YYYY from a string."""
    m = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", s)
    if not m: return None
    try:
        return date_type(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    except ValueError:
        return None


def _normalize_rig(name: str) -> str:
    """'ENF # 08' / 'ENF#08' / 'ENF-08' → 'ENF-08'."""
    s = _clean(name).upper()
    m = re.search(r"ENF[\s#\-]*(\d+)", s)
    if m: return f"ENF-{m.group(1)}"
    return _clean(name)


# ---------------------------------------------------------------------------
# Operations parser
# ---------------------------------------------------------------------------
def _parse_parallel_ops(time_lines: List[str], desc_lines: List[str]) -> List[dict]:
    """
    Time bands and descriptions live in separate parallel cells, with
    lines aligned by index:

        cell 0 ('HORRAIRE'):              cell 1 (descriptions):
        L0: 00h00-12h00                   L0: DTM: suite réception appareil …
        L1: 12h00-16h00                   L1: Traitement de boue à l'huile …
        L2: 16h00-17h00                   L2: Montage ligne de pompage …

    Rule: a line at index i in cell 0 with a parseable HHhMM-HHhMM band
    starts a new op; the same index in cell 1 is its description.  A
    blank / unparseable time-band line at index i means cell 1's line
    at index i is a continuation of the previous op's description.
    """
    activities = []
    current = None
    n = max(len(time_lines), len(desc_lines))
    time_lines = list(time_lines) + [""] * (n - len(time_lines))
    desc_lines = list(desc_lines) + [""] * (n - len(desc_lines))

    for i in range(n):
        tl = time_lines[i].strip()
        dl = desc_lines[i].rstrip()
        m = re.match(r"^(\d{1,2})[hH:](\d{2})\s*[-–]\s*(\d{1,2})[hH:](\d{2})\s*$", tl)
        if m:
            if current is not None:
                activities.append(_finalize_op(current))
            start_t = time(int(m.group(1)) % 24, int(m.group(2)))
            end_t   = time(int(m.group(3)) % 24, int(m.group(4)))
            current = {
                "start_time": start_t,
                "end_time":   end_t,
                "desc_parts": [dl.strip()] if dl.strip() else [],
            }
        else:
            if current is not None and dl.strip():
                current["desc_parts"].append(dl.strip())
    if current is not None:
        activities.append(_finalize_op(current))
    return activities


def _finalize_op(d: dict) -> dict:
    start_t = d["start_time"]
    end_t   = d["end_time"]
    hours   = _hours_between(start_t, end_t)
    return {
        "start_time": start_t,
        "end_time":   end_t,
        "hours":      hours,
        "phase_name": "",
        "code": "", "sub": "",
        "description": _clean(" ".join(d["desc_parts"])),
        "start_md": 0, "end_md": 0,
        "npt": 0, "npt_detail": "",
        "npt_company": "", "op_company": "",
        "bill": "",        # back-filled from tarif totals
    }


# ---------------------------------------------------------------------------
# Cell-by-cell table extraction with merged-cell handling
# ---------------------------------------------------------------------------
def _row_cells(row) -> List[str]:
    """Return one entry per VISUAL cell.  Word's `cells` collection
    repeats the same cell object for spans across columns, so adjacent
    cells with identical text are deduplicated — but only when their
    underlying tc element is the same (true merged cell), not when two
    different cells just happen to hold the same string."""
    out = []
    prev_tc = None
    for c in row.cells:
        tc = c._tc        # underlying XML element
        if tc is prev_tc:
            continue      # same merged cell — skip
        out.append(c.text.strip())
        prev_tc = tc
    return out


def _find_value_after(text_row: List[str], label_substr: str) -> str:
    """In a list of row cells, find the cell whose normalized text EQUALS
    label_substr (case-insensitive, whitespace-normalized) and return the
    next cell's value.  Falls back to "starts with" matching."""
    norm = lambda s: re.sub(r"\s+", " ", s.strip().lower())
    target = norm(label_substr)
    for i, cell in enumerate(text_row):
        if norm(cell) == target and i + 1 < len(text_row):
            return text_row[i + 1].strip()
    for i, cell in enumerate(text_row):
        if norm(cell).startswith(target) and i + 1 < len(text_row):
            return text_row[i + 1].strip()
    return ""


def _is_horaire_header(cells: List[str]) -> bool:
    """The time-band header cell is spelled 'HORAIRE' in most reports of
    this family but 'HORRAIRE' (double-R typo) in others — match both."""
    for c in cells:
        u = c.upper()
        if "HORAIRE" in u or "HORRAIRE" in u:
            return True
    return False


# ---------------------------------------------------------------------------
# Main extractor
# ---------------------------------------------------------------------------
def parse_enf08(source: Union[Path, str, BytesIO]) -> dict:
    """Parse an ENF-08 (Word .doc / .docx) workover report."""
    if isinstance(source, BytesIO):
        try:
            doc = Document(source)
            tmp_path = None
        except Exception:
            with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tf:
                tf.write(source.getvalue())
                tmp_path = Path(tf.name)
            docx_path = _ensure_docx(tmp_path)
            doc = Document(str(docx_path))
    else:
        source_path = Path(source)
        docx_path = _ensure_docx(source_path)
        doc = Document(str(docx_path))

    # Extract paragraphs and tables
    paragraphs = [_clean(p.text) for p in doc.paragraphs]
    tables = doc.tables

    # =====================================================================
    # HEADER (table 0 + paragraphs)
    # =====================================================================
    header = {}

    # Date + report number — from paragraph "RAPPORT N° 01 du 19/07/2026"
    for p in paragraphs:
        m = re.search(r"du\s+(\d{1,2}[/-]\d{1,2}[/-]\d{4})", p, re.IGNORECASE)
        if m:
            d = _date_from_text(m.group(1))
            if d:
                header["date"] = d
                break
    for p in paragraphs:
        m = re.search(r"RAPPORT\s*N\s*[°o]\s*(\d+)", p, re.IGNORECASE)
        if m:
            header["day_number"] = int(m.group(1))
            break

    # Fallback: "I/ APPAREIL N° : ENF#08   PUITS : ISNO#02   RAPPORT: N° 01 du 19/07/2026"
    for p in paragraphs:
        if "APPAREIL" in p.upper() and "PUITS" in p.upper():
            m = re.search(r"APPAREIL\s*N\s*[°o]?\s*:?\s*([A-Z0-9#\-\s]+?)\s{2,}", p, re.IGNORECASE)
            if m and "rig_name" not in header:
                header["rig_name"] = _normalize_rig(m.group(1))
            m = re.search(r"PUITS\s*:?\s*([A-Z0-9#\-]+)", p, re.IGNORECASE)
            if m and "well_name" not in header:
                header["well_name"] = m.group(1).replace(" ", "")
            if "date" not in header:
                m = re.search(r"du\s+(\d{1,2}[/-]\d{1,2}[/-]\d{4})", p, re.IGNORECASE)
                if m:
                    d = _date_from_text(m.group(1))
                    if d: header["date"] = d
            break

    # Pull rig/well/field/tool/mud info from table 0
    if tables:
        t0 = tables[0]
        for row in t0.rows:
            cells = _row_cells(row)

            puits = _find_value_after(cells, "puits")
            if puits and "well_name" not in header:
                header["well_name"] = puits.replace(" ", "")
            champ = _find_value_after(cells, "champ")
            if champ and "field_name" not in header:
                header["field_name"] = champ
            app = _find_value_after(cells, "appareil")
            if app and "rig_name" not in header:
                header["rig_name"] = _normalize_rig(app)

            # Tubing / completion tool designation (replaces TP-186's BHA/bit fields)
            desig = _find_value_after(cells, "désignation")
            if desig: header["_tool_designation"] = desig
            tdiam = _find_value_after(cells, "diamètre")
            if tdiam: header["_tool_diameter"] = tdiam
            marque = _find_value_after(cells, "marque-type")
            if marque: header["_tool_brand"] = marque

            # Mud volumes
            vt = _find_value_after(cells, "volume total")
            if vt: header["_mud_total"] = vt
            vs = _find_value_after(cells, "volume surface")
            if vs: header["_mud_surface"] = vs
            vp = _find_value_after(cells, "volume puits")
            if vp: header["_mud_well"] = vp

            # Mud density / viscosity / type
            for label, key in [("densité", "_density"), ("v. marsh", "_fun_vis"),
                                ("type", "_mud_type")]:
                v = _find_value_after(cells, label)
                if v:
                    header[key] = v

            # Filtrat here is a Huile/Eau (oil/water) split ratio, e.g. "90/10",
            # not a single filtrate number — capture the label's *second*
            # follow-on cell (the ratio), and the descriptive one separately.
            for i, cell in enumerate(cells):
                if re.sub(r"\s+", " ", cell.strip().lower()).startswith("filtrat"):
                    if i + 1 < len(cells):
                        header["_oil_water_label"] = cells[i + 1]
                    if i + 2 < len(cells):
                        header["_oil_water_ratio"] = cells[i + 2]
                    break

            # Composition garniture (DP / DC)
            dp = _find_value_after(cells, "dp (diam / nbre)")
            if dp: header["_dp_composition"] = dp
            dc = _find_value_after(cells, "dc (diam / nbre)")
            if dc: header["_dc_composition"] = dc

    # Supervisor — paragraph after "Responsable de la Section"
    for i, p in enumerate(paragraphs):
        if "RESPONSABLE" in p.upper() and "SECTION" in p.upper():
            for j in range(i + 1, min(i + 6, len(paragraphs))):
                cand = paragraphs[j]
                if cand and len(cand) < 60 and not cand.startswith(("S", "RAPPORT", "OPERATION")):
                    if re.match(r"^[A-Z][\.\s]", cand) or "." in cand[:5]:
                        header["supervisor"] = cand
                        break
            break

    # =====================================================================
    # OPERATIONS (table 1)
    # Layout: a single row with two parallel cells — cell 0 holds time
    # bands ('HHhMM-HHhMM' newline-separated), cell 1 holds the matching
    # descriptions (also newline-separated), aligned by line index. The
    # header row spells this "HORAIRE" or (typo) "HORRAIRE".
    # =====================================================================
    activities = []
    if len(tables) > 1:
        t1 = tables[1]
        for ri, row in enumerate(t1.rows):
            cells = _row_cells(row)
            if _is_horaire_header(cells):
                if ri + 1 < len(t1.rows):
                    data_cells = _row_cells(t1.rows[ri + 1])
                    if len(data_cells) >= 2:
                        time_lines = data_cells[0].split("\n")
                        desc_lines = data_cells[1].split("\n")
                        activities = _parse_parallel_ops(time_lines, desc_lines)
                break

    # =====================================================================
    # TARIF TOTALS (table 2) — rows labelled "T. 1" / "T. 2" / "T. 3" / "T. 4"
    # (dot+space variant of TP-186's "T1"/"T2"); "N/R" rows are ignored.
    # =====================================================================
    tarif_totals = {}
    if len(tables) > 2:
        t2 = tables[2]
        for row in t2.rows:
            cells = _row_cells(row)
            if len(cells) < 2: continue
            label = _clean(cells[0]).upper().replace(" ", "").replace(".", "")
            if label in ("T1", "T2", "T3", "T4"):
                hrs = _parse_hhmm_duration(cells[1])
                if hrs > 0:
                    tarif_totals[label.lower()] = hrs

    # Back-fill bill codes via the shared helper
    from helpers.bill_code_assign import assign_bill_codes
    activities = assign_bill_codes(activities, tarif_totals)

    # =====================================================================
    # COSTS
    # Preferred source: explicit "Coût journalier" cell in table 1.
    # Fallback: table 2 TOTAL row's rightmost monetary value (TP-186 style).
    # =====================================================================
    if len(tables) > 1:
        for row in tables[1].rows:
            cells = _row_cells(row)
            if cells and "COÛT JOURNALIER" in cells[0].upper():
                if len(cells) > 1:
                    v = _money_to_int(cells[1])
                    if v > 0:
                        header["daily_cost"] = v
                break

    if "daily_cost" not in header and len(tables) > 2:
        t2 = tables[2]
        for row in t2.rows:
            cells = _row_cells(row)
            if cells and "TOTAL" in cells[0].upper():
                for c in reversed(cells):
                    if "DA" in c.upper():
                        v = _money_to_int(c)
                        if v > 0:
                            header["daily_cost"] = v
                            break
                break

    # Cumul depuis origine — paragraph
    for p in paragraphs:
        if "CUMUL DEPUIS ORIGINE" in p.upper() and "APPAREIL" in p.upper():
            v = _money_to_int(p)
            if v > 0:
                header["cum_cost"] = v
                break

    # Cumul depuis origine — autres sociétés (service companies), if present
    for p in paragraphs:
        if "CUMUL DEPUIS ORIGINE" in p.upper() and "SOCIETES" in p.upper():
            v = _money_to_int(p)
            if v > 0:
                header["cum_cost_other_companies"] = v
            break

    # =====================================================================
    # MUD CHECKS + VOLUMES + TOOL/COMPOSITION (parsed from stashed header values)
    # =====================================================================
    mud_checks = {}
    mud_volume = {}
    for src_key, dst_key in [("_density", "density"), ("_fun_vis", "fun_vis")]:
        raw = header.pop(src_key, None)
        if raw is not None:
            v = _float(raw)
            if v: mud_checks[dst_key] = v

    mt = header.pop("_mud_type", None)
    if mt and mt.lower() in ("huile", "obm", "oil", "wbm", "water"):
        mud_checks["mud_type"] = mt

    # Oil/water filtrat split ratio — kept as text, not forced into a float
    ow_label = header.pop("_oil_water_label", None)
    ow_ratio = header.pop("_oil_water_ratio", None)
    if ow_ratio:
        mud_checks["oil_water_ratio"] = ow_ratio
        if ow_label:
            mud_checks["oil_water_label"] = ow_label

    for src_key, dst_key in [("_mud_total", "total_volume"),
                              ("_mud_surface", "pits_volume"),
                              ("_mud_well", "string_volume")]:
        raw = header.pop(src_key, None)
        if raw is not None:
            v = _float(raw)
            if v: mud_volume[dst_key] = v

    # Tool / completion string + garniture composition
    tool_info = {}
    for src_key, dst_key in [("_tool_designation", "designation"),
                              ("_tool_diameter", "diameter"),
                              ("_tool_brand", "brand"),
                              ("_dp_composition", "dp_composition"),
                              ("_dc_composition", "dc_composition")]:
        raw = header.pop(src_key, None)
        if raw: tool_info[dst_key] = _clean(raw)
    if tool_info:
        header["tool_info"] = tool_info

    # Strip any remaining temporary keys
    for k in list(header.keys()):
        if k.startswith("_"):
            header.pop(k, None)

    # =====================================================================
    # MUD CHEMICAL USAGE (table 1: a sub-row after "DESIGNATIONS")
    # =====================================================================
    chemicals = []
    if len(tables) > 1:
        t1 = tables[1]
        for ri, row in enumerate(t1.rows):
            cells = _row_cells(row)
            if any("DESIGNATION" in c.upper() for c in cells):
                if ri + 1 < len(t1.rows):
                    next_cells = _row_cells(t1.rows[ri + 1])
                    if len(next_cells) >= 3:
                        items = next_cells[0].split("\n")
                        useds  = next_cells[1].split("\n") if next_cells[1] else []
                        stocks = next_cells[2].split("\n") if next_cells[2] else []
                        for i, item in enumerate(items):
                            item = item.strip()
                            if not item: continue
                            chemicals.append({
                                "item": item,
                                "units": "",
                                "received": "",
                                "used":   (useds[i].strip() if i < len(useds) else ""),
                                "on_loc": (stocks[i].strip() if i < len(stocks) else ""),
                            })
                break

    # =====================================================================
    # TEXT SECTIONS (table 1: "Situation à 06h00" and "Programme prévu")
    # =====================================================================
    text_sections = {}
    if len(tables) > 1:
        for row in tables[1].rows:
            cells = _row_cells(row)
            if not cells: continue
            label = cells[0].upper() if cells[0] else ""
            value = cells[1] if len(cells) > 1 else ""
            value_clean = _clean(value)
            if not value_clean: continue
            if "SITUATION" in label:
                if len(value_clean) > 300:
                    value_clean = value_clean[:300].rsplit(None, 1)[0] + "…"
                text_sections["current_operation"] = value_clean
                text_sections["day_summary"]       = value_clean
            elif "PROGRAMME" in label:
                text_sections["plan_operations"] = value_clean

    # =====================================================================
    # SERVICE COMPANIES (table 3)
    # =====================================================================
    service_companies = []
    if len(tables) > 3:
        for row in tables[3].rows:
            cells = _row_cells(row)
            if len(cells) < 3: continue
            name = cells[0].strip()
            if not name or name.upper() in ("SOCIETES", "TOTAL"): continue
            if name in ("/",): continue
            service_companies.append({
                "company": name,
                "nature":  cells[1].strip(),
                "amount":  _money_to_int(cells[2]),
            })

    return {
        "header": header,
        "activities": activities,
        "text_sections": text_sections,
        "mud_checks": mud_checks,
        "mud_volume": mud_volume,
        "mud_chemical_usage": chemicals,
        "personnel_data": [],
        "pumps": [],
        "well_location": {},
        "survey_data": [],
        "safety": {},
        "tarif_totals": tarif_totals,
        "service_companies": service_companies,
    }


# Drop-in compat
parse_daily_excel_report = parse_enf08


if __name__ == "__main__":
    import sys, json
    if len(sys.argv) < 2:
        sys.exit("Usage: enf08_extract.py SOURCE.doc|SOURCE.docx")
    data = parse_enf08(Path(sys.argv[1]))

    def default(o):
        if isinstance(o, (date_type, datetime)): return o.isoformat()
        if isinstance(o, time): return o.strftime("%H:%M:%S")
        if isinstance(o, timedelta): return o.total_seconds()
        return str(o)
    print(json.dumps(data, indent=2, default=default, ensure_ascii=False))
